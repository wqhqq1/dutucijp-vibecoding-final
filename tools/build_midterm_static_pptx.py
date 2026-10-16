"""Build the static backup PPTX for the midterm talk from the web deck.

The web deck (中期汇报网页版/index.html) is the master. Each slide is rendered by
headless Chrome at 2x and placed as one full-slide picture; speaker notes are the
plain text of the slide's <aside class="notes">. The PPTX is not editable text on
purpose: it is only a playback backup with the same content as the web/PDF version.

Usage (python-pptx is installed into the ignored .working/pydeps, not the project):
  python -I tools/build_midterm_static_pptx.py
"""
from __future__ import annotations
import hashlib, html, json, os, re, subprocess, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.working' / 'pydeps'))
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

DECK = ROOT / '中期汇报网页版' / 'index.html'
OUT = ROOT / '中期汇报.pptx'
SHOTS = ROOT / '.working' / 'webdeck' / 'hi'
CHROME = Path(os.environ.get('ProgramFiles', 'C:' + os.sep + 'Program Files')) / 'Google' / 'Chrome' / 'Application' / 'chrome.exe'


def notes_of(source: str) -> list[str]:
    result = []
    for sec in re.findall(r'<section class="slide.*?</section>', source, re.S):
        m = re.search(r'<aside class="notes">(.*?)</aside>', sec, re.S)
        parts = re.findall(r'<p[^>]*>(.*?)</p>', m.group(1), re.S) if m else []
        text = [html.unescape(re.sub(r'<[^>]+>', '', p)).strip() for p in parts]
        title = re.search(r'data-title="([^"]+)"', sec).group(1)
        sec_budget = re.search(r'data-sec="(\d+)"', sec).group(1)
        lines = [f'{title} · 建议 {sec_budget} 秒', text[0] if text else '']
        lines += ['备查：' + t for t in text[1:]]
        result.append('\n'.join(lines))
    return result


def main() -> None:
    source = DECK.read_text(encoding='utf-8')
    notes = notes_of(source)
    SHOTS.mkdir(parents=True, exist_ok=True)
    for i in range(len(notes)):
        shot = SHOTS / f's{i}.png'
        shot.unlink(missing_ok=True)
        subprocess.run([str(CHROME), '--headless=new', '--disable-gpu', '--hide-scrollbars',
                        '--window-size=1600,900', '--force-device-scale-factor=2',
                        '--virtual-time-budget=2500', f'--screenshot={shot}',
                        DECK.as_uri() + f'?static#{i}'], check=True, capture_output=True)
        if not shot.exists():
            raise SystemExit(f'render failed: slide {i}')
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for i, text in enumerate(notes):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(SHOTS / f's{i}.png'), 0, 0, prs.slide_width, prs.slide_height)
        slide.notes_slide.notes_text_frame.text = text
    prs.core_properties.title = '中期汇报（v1.2）· 退货对账与差异处理系统 · 静态备份'
    prs.core_properties.author = '王麒皓 20251101030'
    prs.save(OUT)
    digest = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    print(json.dumps({'pptx': str(OUT.relative_to(ROOT)), 'slides': len(notes),
                      'pptx_sha256': digest(OUT), 'deck_sha256': digest(DECK)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
