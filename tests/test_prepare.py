"""Isolated fixtures for cleaning rules; never substitutes for official business data."""
import csv
import gc
import json
import sys
import tempfile
import unittest
from zipfile import ZipFile
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'tools'))
from prepare import HEADERS, SHEETS, prepare, raw_rows
from fetch_data import sha256


class CleaningTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)

    def fixture(self, first, second=()):
        book = openpyxl.Workbook()
        book.remove(book.active)
        for title, rows in zip(SHEETS, (first, second)):
            sheet = book.create_sheet(title)
            sheet.append(HEADERS)
            for row in rows:
                sheet.append(row)
        path = self.root/'fixture.xlsx'
        book.save(path)
        return path

    def sale(self):
        return ['123', 'ABC-1', 'Fixture only', 10, datetime(2010,1,1,9), 0.1, 42, 'UK']

    def run_prepare(self, first, second=(), scope='all'):
        path = self.fixture(first, second)
        before = sha256(path)
        target = prepare(path, scope, self.root/'prepared', progress=lambda _:None)
        self.assertEqual(before, sha256(path))
        with (target/'records.csv').open(encoding='utf-8',newline='') as stream:
            rows = list(csv.DictReader(stream))
        manifest=json.loads((target/'manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(manifest['stats']['raw'], manifest['stats']['accepted']+manifest['stats']['excluded'])
        return rows,manifest

    def test_TC03_header_and_alphanumeric_stock(self):
        rows,m=self.run_prepare([self.sale()]);self.assertEqual(rows[0]['stock'],'ABC-1');self.assertEqual(rows[0]['customer'],'42');self.assertEqual(rows[0]['price'],'0.1');self.assertEqual(rows[0]['source_row'],'2')

    def test_TC03_text_identifiers_preserve_leading_zeroes(self):
        row=self.sale();row[0]='00123';row[1]='01A';row[6]='00042'
        rows,_=self.run_prepare([row])
        self.assertEqual((rows[0]['invoice'],rows[0]['stock'],rows[0]['customer']),('00123','01A','00042'))

    def test_TC03_exact_XML_price_survives_beyond_binary_float_precision(self):
        path=self.fixture([self.sale()])
        exact='0.1000000000000000000000000001'
        with ZipFile(path) as source:
            parts={name:source.read(name) for name in source.namelist()}
        parts['xl/worksheets/sheet1.xml']=parts['xl/worksheets/sheet1.xml'].replace(b'<v>0.1</v>', ('<v>'+exact+'</v>').encode())
        with ZipFile(path,'w') as target:
            for name,content in parts.items():target.writestr(name,content)
        prepared=prepare(path,'year1',self.root/'prepared',progress=lambda _:None)
        with (prepared/'records.csv').open(encoding='utf-8',newline='') as stream:
            row=next(csv.DictReader(stream))
        self.assertEqual(row['price'],exact)
        self.assertEqual(json.loads(row['raw_json'])[5],['n',exact])

    def test_TC04_invalid_price_quantity_and_case_insensitive_C(self):
        values=[]
        for invoice,quantity,price in [('C1',4,1),('123',-4,1),('C2',-4,0),('123',1.5,1),('c3',-4,1)]:
            r=self.sale();r[0],r[3],r[5]=invoice,quantity,price;values.append(r)
        rows,m=self.run_prepare(values)
        self.assertEqual([r['excluded'] for r in rows],['invalid_quantity','invalid_quantity','nonpositive_price','invalid_quantity',''])
        self.assertEqual(rows[-1]['kind'],'cancel')

    def test_TC05_missing_identity_retained_description_uninvented(self):
        r=self.sale();r[0]='C1';r[3]=-4;r[2]=r[6]=r[7]=None
        rows,m=self.run_prepare([r]);r=rows[0];self.assertEqual(r['excluded'],'');self.assertEqual(r['description'],'');self.assertEqual(r['flags'],'missing_customer|missing_country|missing_description')

    def test_TC06_cross_sheet_occurrence_alignment_2_and_3(self):
        a=self.sale();rows,m=self.run_prepare([a,a],[a,a,a]);self.assertEqual([r['excluded'] for r in rows],['','','cross_sheet_duplicate','cross_sheet_duplicate','']);self.assertEqual([r['duplicate_of'] for r in rows[2:]],['0:2','0:3','']);self.assertEqual(m['stats']['raw'],5)

    def test_TC07_same_sheet_duplicates_all_marked_not_deleted(self):
        a=self.sale();rows,m=self.run_prepare([a,a]);self.assertEqual(m['stats']['accepted'],2);self.assertTrue(all('same_sheet_duplicate' in r['flags'] for r in rows))

    def test_TC08_exclusion_priority_and_single_main_reason(self):
        a=self.sale();a[0]=None;a[3]=-1;a[5]=0
        rows,m=self.run_prepare([a],[a]);self.assertEqual(rows[0]['excluded'],'invalid_field');self.assertEqual(rows[1]['excluded'],'cross_sheet_duplicate');self.assertEqual(sum(m['stats']['exclusion_reasons'].values()),2)

    def test_TC02_scope_selected_before_cleaning(self):
        rows,m=self.run_prepare([self.sale()]*10002,[self.sale()],scope='benchmark');self.assertEqual(m['stats']['raw'],10000);self.assertEqual(rows[0]['source_row'],'2');self.assertEqual(rows[-1]['source_row'],'10001')

    def test_prepared_hash_tampering_rejected(self):
        path=self.fixture([self.sale()]);target=prepare(path,'year1',self.root/'prepared',progress=lambda _:None)
        with (target/'records.csv').open('a',encoding='utf-8') as stream:stream.write('tamper')
        with self.assertRaisesRegex(ValueError,'hash mismatch'):prepare(path,'year1',self.root/'prepared',progress=lambda _:None)

    def test_source_1904_calendar_and_nonstandard_cell_types_remain_exact(self):
        first=self.sale();first[0]='=1+1';first[3]=True
        path=self.fixture([first,self.sale()])
        book=openpyxl.load_workbook(path)
        book.epoch=openpyxl.utils.datetime.CALENDAR_MAC_1904
        book.save(path);book.close()
        source=list(raw_rows(path,'year1'))
        self.assertEqual(source[0][3][0],('formula','1+1'))
        self.assertEqual(source[0][3][3],('b','1'))
        target=prepare(path,'year1',self.root/'prepared',progress=lambda _:None)
        with (target/'records.csv').open(encoding='utf-8',newline='') as stream:
            rows=list(csv.DictReader(stream))
        self.assertEqual(rows[0]['excluded'],'invalid_field')
        self.assertEqual(rows[1]['time'],'2010-01-01T09:00:00.000000')
        self.assertEqual(json.loads(rows[0]['raw_json'])[0],['formula','1+1'])

    def test_preparation_restores_cycle_collection_after_success_and_failure(self):
        enabled=gc.isenabled()
        self.addCleanup(gc.enable if enabled else gc.disable)
        path=self.fixture([self.sale()])
        for initial in (True,False):
            (gc.enable if initial else gc.disable)()
            prepare(path,'year1',self.root/f'prepared-{initial}',progress=lambda _:None)
            self.assertEqual(gc.isenabled(),initial)
            with self.assertRaises(FileNotFoundError):
                prepare(self.root/'missing.xlsx','all',self.root/'failed',progress=lambda _:None)
            self.assertEqual(gc.isenabled(),initial)

    def test_shared_and_inline_rich_text_and_empty_numeric_lexemes_remain_exact(self):
        path=self.fixture([self.sale()])
        with ZipFile(path) as source:
            parts={name:source.read(name) for name in source.namelist()}
        xml=parts['xl/worksheets/sheet1.xml']
        start=xml.index(b'<c r="A2"');end=xml.index(b'</c>',start)+4
        xml=xml[:start]+b'<c r="A2" t="s"><v>0</v></c>'+xml[end:]
        start=xml.index(b'<c r="C2"');end=xml.index(b'</c>',start)+4
        xml=xml[:start]+b'<c r="C2" t="inlineStr"><is><r><t> A&amp;</t></r><r><t> B </t></r></is></c>'+xml[end:]
        start=xml.index(b'<c r="F2"');end=xml.index(b'</c>',start)+4
        xml=xml[:start]+b'<c r="F2" t="n"><v/></c>'+xml[end:]
        parts['xl/worksheets/sheet1.xml']=xml
        parts['xl/sharedStrings.xml']=b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><t>00</t></r><r><t>123</t></r></si></sst>'
        with ZipFile(path,'w') as target:
            for name,content in parts.items():target.writestr(name,content)
        source=list(raw_rows(path,'year1'))
        self.assertEqual(source[0][3][0],('s','00123'))
        self.assertEqual(source[0][3][2],('s',' A& B '))
        self.assertEqual(source[0][3][5],('n',''))
        target=prepare(path,'year1',self.root/'prepared',progress=lambda _:None)
        with (target/'records.csv').open(encoding='utf-8',newline='') as stream:row=next(csv.DictReader(stream))
        self.assertEqual(row['invoice'],'00123')
        self.assertEqual(row['description'],'A& B')
        self.assertEqual(row['excluded'],'invalid_field')

    def test_unexpected_workbook_sheet_metadata_is_rejected(self):
        path=self.fixture([self.sale()])
        book=openpyxl.load_workbook(path);book.worksheets[1].title='Unexpected sheet';book.save(path);book.close()
        with self.assertRaisesRegex(ValueError,'Unexpected worksheets'):
            list(raw_rows(path,'year1'))


if __name__ == '__main__':
    unittest.main(verbosity=2)
