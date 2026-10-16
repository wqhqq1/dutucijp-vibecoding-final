"""CR-01..06. Exact XML price lexemes; openpyxl handles Excel calendar metadata.

Produces immutable CSV + manifest outside the repository. All raw eight fields are
retained. Decimal display rounding is never used to form an identity or match key.
"""
import argparse
import csv
import gc
import hashlib
import json
import os
import tempfile
import time
from collections import Counter
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.datetime import from_excel

from fetch_data import sha256

CLEANING = 'cr-1.0.1'
SHEETS = ('Year 2009-2010', 'Year 2010-2011')
HEADERS = ('Invoice', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'Price', 'Customer ID', 'Country')
FIELDS = ('id', 'sheet_index', 'sheet', 'source_row', 'raw_json', 'invoice', 'stock', 'description',
          'quantity', 'time', 'price', 'customer', 'country', 'kind', 'excluded', 'flags', 'duplicate_of')
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
ROW_TAG, VALUE_TAG, FORMULA_TAG, TEXT_TAG, STRING_TAG = (NS+tag for tag in ('row', 'v', 'f', 't', 'si'))
encode_raw = json.JSONEncoder(ensure_ascii=False, separators=(',', ':')).encode


@lru_cache(maxsize=65536)
def canonical(value):
    number = Decimal(value)
    if not number.is_finite():
        raise ValueError('nonfinite decimal')
    if not number:
        return '0'
    formatted = format(number, 'f')
    return formatted.rstrip('0').rstrip('.') if '.' in formatted else formatted


@lru_cache(maxsize=65536)
def source_time(value, epoch):
    """Memoize the same Excel conversion; repeated invoice dates are frequent."""
    rendered = from_excel(float(value), epoch).isoformat(timespec='microseconds')
    if 'T' not in rendered:
        raise ValueError('invalid date')
    return rendered


@lru_cache(maxsize=65536)
def integer(value):
    number = Decimal(value)
    if not number.is_finite() or number != number.to_integral_value() or abs(number) > 2**63-1:
        raise ValueError('invalid integer')
    return int(number)


@lru_cache(maxsize=131072)
def raw_cell(typ, value):
    """Share immutable repeated XML values without changing their source lexeme."""
    return (typ, value)


def raw_rows(path, scope):
    """Decode original cell types/values; don't round numeric prices to float."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if tuple(book.sheetnames) != SHEETS:
            raise ValueError('Unexpected worksheets: '+repr(book.sheetnames))
        epoch = book.epoch
    finally:
        book.close()
    with ZipFile(path) as zipped:
        strings = []
        if 'xl/sharedStrings.xml' in zipped.namelist():
            with zipped.open('xl/sharedStrings.xml') as stream:
                for _, elem in ET.iterparse(stream, events=('end',)):
                    if elem.tag == STRING_TAG:
                        strings.append(('s', ''.join(t.text or '' for t in elem.iter(TEXT_TAG))))
                        elem.clear()
        for si, sheet in enumerate(SHEETS[:2 if scope == 'all' else 1]):
            with zipped.open(f'xl/worksheets/sheet{si+1}.xml') as stream:
                # Only row-end events are consumed; subscribing to starts also
                # allocated/dispatched millions of unused events per workbook.
                parser = ET.iterparse(stream, events=('end',))
                for _, row in parser:
                    if row.tag != ROW_TAG:
                        continue
                    rn = int(row.attrib['r'])
                    if scope == 'benchmark' and rn > 10001:
                        break
                    values = [None]*8
                    for cell in row:
                        column = cell.attrib['r'].rstrip('0123456789')
                        if len(column) != 1 or not 'A' <= column <= 'H':
                            raise ValueError('Unexpected column '+column)
                        typ = cell.get('t', 'n')
                        value = cell.findtext(VALUE_TAG)
                        if cell.find(FORMULA_TAG) is not None:
                            values[ord(column)-65] = ('formula', cell.findtext(FORMULA_TAG))
                        elif typ == 's':
                            values[ord(column)-65] = strings[int(value)]
                        elif typ == 'inlineStr':
                            values[ord(column)-65] = raw_cell('s', ''.join(t.text or '' for t in cell.iter(TEXT_TAG)))
                        elif value is not None:
                            values[ord(column)-65] = raw_cell(typ, value)
                    if rn == 1:
                        if tuple(v[1] if v else None for v in values) != HEADERS:
                            raise ValueError('Unexpected source headers')
                    else:
                        yield si, sheet, rn, tuple(values), epoch
                    row.clear()


def text(cell):
    if cell is None:
        return ''
    typ, value = cell
    if typ == 's':
        return value.strip()
    if typ == 'n':
        return canonical(value)
    raise ValueError('Invalid identity type')


def fingerprint(raw):
    # Numeric equivalents have the same field value; text and numeric types stay distinct.
    normalized = []
    for cell in raw:
        if cell is None:
            normalized.append(None)
            continue
        typ, value = cell
        if typ == 'n':
            try:
                value = canonical(value)
            except (ValueError, InvalidOperation):
                pass  # A bad numeric cell is classified by cleaning, not silently dropped.
        normalized.append(raw_cell(typ, value))
    # This key is internal, never persisted. Tuple equality preserves all eight
    # typed values without serializing/hashing an extra JSON document per row.
    return tuple(normalized)


def normalize(si, sheet, rn, raw, epoch):
    row = dict.fromkeys(FIELDS, '')
    row.update(id=f'{si}:{rn}', sheet_index=si, sheet=sheet, source_row=rn,
               raw_json=encode_raw(raw))
    flags = []
    problems = []
    try:
        row['invoice'], row['stock'] = text(raw[0]), text(raw[1])
        if not row['invoice'] or not row['stock']:
            raise ValueError('missing invoice or stock')
        row['kind'] = 'cancel' if row['invoice'].upper().startswith('C') else 'sale'
        row['description'] = text(raw[2])
        row['country'] = text(raw[7])
        row['customer'] = text(raw[6])
        if not row['customer']:
            flags.append('missing_customer')
        if not row['country']:
            flags.append('missing_country')
        if not row['description']:
            flags.append('missing_description')
        if not raw[4] or raw[4][0] != 'n':
            raise ValueError('invalid date type')
        row['time'] = source_time(raw[4][1], epoch)
        if not raw[5] or raw[5][0] != 'n':
            raise ValueError('invalid price type')
        row['price'] = canonical(raw[5][1])
        if Decimal(row['price']) <= 0:
            problems.append('nonpositive_price')
        if not raw[3] or raw[3][0] != 'n':
            problems.append('invalid_quantity')
        else:
            try:
                row['quantity'] = integer(raw[3][1])
                if (row['kind'] == 'sale' and row['quantity'] <= 0) or (row['kind'] == 'cancel' and row['quantity'] >= 0):
                    problems.append('invalid_quantity')
            except (ValueError, InvalidOperation):
                problems.append('invalid_quantity')
    except (ValueError, InvalidOperation, OverflowError, AttributeError):
        problems.insert(0, 'invalid_field')
    row['excluded'] = problems[0] if problems else ''
    row['flags'] = flags
    return row


def prepare(path, scope, output_root, progress=print):
    # The XML/row graph is acyclic and is released by reference counting. Avoid
    # repeatedly scanning a growing million-row graph for cycles while loading.
    # Always restore the caller's setting, including when parsing fails.
    collect_cycles = gc.isenabled()
    try:
        gc.disable()
        return _prepare(path, scope, output_root, progress)
    finally:
        if collect_cycles:
            gc.enable()


def _prepare(path, scope, output_root, progress):
    started = time.perf_counter()
    source_hash = sha256(path)
    dataset_id = hashlib.sha256(f'{source_hash}|{scope}|{CLEANING}'.encode()).hexdigest()
    target = output_root/dataset_id
    if target.exists():
        manifest = json.loads((target/'manifest.json').read_text(encoding='utf-8'))
        if manifest['records_sha256'] != sha256(target/'records.csv'):
            raise ValueError('Prepared data hash mismatch; preserve it for investigation')
        progress(json.dumps(dict(stage='reused_prepared', dataset_id=dataset_id, path=str(target))))
        return target
    output_root.mkdir(parents=True, exist_ok=True)
    scratch = Path(tempfile.mkdtemp(prefix='preparing-', dir=output_root))
    signatures = {}
    rows = []
    for si, sheet, rn, raw, epoch in raw_rows(path, scope):
        fp = fingerprint(raw)
        group = signatures.get(fp)
        if group is None:
            group = [0, 0, []]
            signatures[fp] = group
        occurrence = group[si]
        group[si] += 1
        duplicate_of = ''
        if si == 0:
            group[2].append(f'{si}:{rn}')
        elif occurrence < len(group[2]):
            duplicate_of = group[2][occurrence]
        rows.append((si, sheet, rn, raw, epoch, group, duplicate_of))
        if len(rows) % 50000 == 0:
            progress(json.dumps(dict(stage='read_source', rows=len(rows))))
    excluded = Counter()
    flags_count = Counter()
    categories = Counter()
    raw_cancel = excluded_cancel = unknown_invoice = 0
    with (scratch/'records.csv').open('w', newline='', encoding='utf-8') as stream:
        writer = csv.writer(stream)
        writer.writerow(FIELDS)
        for si, sheet, rn, raw, epoch, group, duplicate_of in rows:
            row = normalize(si, sheet, rn, raw, epoch)
            if group[si] > 1:
                row['flags'].append('same_sheet_duplicate')
            if duplicate_of:
                row['excluded'] = 'cross_sheet_duplicate'
                row['duplicate_of'] = duplicate_of
            invoice = raw[0]
            invoice_text = text(invoice) if invoice and invoice[0] in ('s','n') else ''
            is_cancel = invoice_text.upper().startswith('C')
            raw_cancel += is_cancel
            unknown_invoice += not bool(invoice_text)
            if row['excluded']:
                excluded[row['excluded']] += 1
                excluded_cancel += is_cancel
            else:
                categories[row['kind']] += 1
            flags_count.update(row['flags'])
            row['flags'] = '|'.join(row['flags'])
            # normalize initializes keys in FIELDS order; updating values does
            # not reorder them. Avoid DictWriter's per-column Python generator.
            writer.writerow(row.values())
    if source_hash != sha256(path):
        raise ValueError('Original changed during preparation; nothing published')
    stats = dict(raw=len(rows), accepted=sum(categories.values()), excluded=sum(excluded.values()),
                 exclusion_reasons=dict(excluded), flags=dict(flags_count), accepted_sales=categories['sale'],
                 accepted_cancels=categories['cancel'], raw_cancels=raw_cancel, excluded_cancels=excluded_cancel,
                 unknown_invoice=unknown_invoice)
    assert stats['raw'] == stats['accepted'] + stats['excluded']
    assert raw_cancel == categories['cancel'] + excluded_cancel
    receipt_path = path.parent/'source.json'
    receipt = json.loads(receipt_path.read_text(encoding='utf-8')) if receipt_path.exists() else {'source_url':'test fixture', 'license':'test only'}
    manifest = dict(dataset_id=dataset_id, scope=scope, cleaning_version=CLEANING, source_hash=source_hash,
                    source=receipt, sheets=list(SHEETS[:2 if scope=='all' else 1]), stats=stats,
                    records_sha256=sha256(scratch/'records.csv'), seconds=time.perf_counter()-started,
                    source_time_semantics='timezone unspecified in source', status='prepared')
    (scratch/'manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False)+'\n', encoding='utf-8')
    scratch.rename(target)
    progress(json.dumps(dict(stage='prepared', path=str(target), **manifest), ensure_ascii=False))
    return target


def main():
    parser = argparse.ArgumentParser()
    root = Path(os.environ['LOCALAPPDATA'])/'dutucijp-vibecoding-final'
    parser.add_argument('--source', type=Path, default=root/'raw/online_retail_II.xlsx')
    parser.add_argument('--scope', choices=('benchmark','year1','all'), required=True)
    parser.add_argument('--output', type=Path, default=root/'prepared')
    args = parser.parse_args()
    prepare(args.source, args.scope, args.output)


if __name__ == '__main__':
    main()
